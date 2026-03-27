# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule

XLSX_PATH = r"F:\cloneSwarm\HeartbeatGraph_SwarmClone.xlsx"

# ── Colors ──────────────────────────────────────────────────────────────────
DARK_BLUE     = "1F3864"
WHITE         = "FFFFFF"
LIGHT_YELLOW  = "FFFFC0"
BLUE_TEXT     = "0000FF"
GREEN_BG      = "C6EFCE"
GRAY_TEXT     = "808080"
ALT_ROW_GRAY  = "F0F0F0"
ORANGE        = "FF8C00"
RED           = "FF0000"
GREEN_TEXT    = "008000"

def dark_fill():
    return PatternFill("solid", fgColor=DARK_BLUE)

def white_bold(size=11, color=WHITE):
    return Font(name="Arial", bold=True, size=size, color=color)

def plain_font(bold=False, color="000000", italic=False, size=11):
    return Font(name="Arial", bold=bold, color=color, italic=italic, size=size)

def center_align(wrap=False):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

def left_align(wrap=False):
    return Alignment(horizontal="left", vertical="center", wrap_text=wrap)

# ────────────────────────────────────────────────────────────────────────────

wb = openpyxl.load_workbook(XLSX_PATH)

# ============================================================
# SHEET 4: "Level Up Estimate"
# ============================================================
if "Level Up Estimate" in wb.sheetnames:
    del wb["Level Up Estimate"]

ws4 = wb.create_sheet("Level Up Estimate")

# ── Helper: write dark-blue header ──────────────────────────
def write_section_header(ws, row, first_col, last_col, text):
    ws.merge_cells(start_row=row, start_column=first_col,
                   end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=first_col, value=text)
    cell.fill  = dark_fill()
    cell.font  = Font(name="Arial", bold=True, size=12, color=WHITE)
    cell.alignment = center_align()

def write_header_row(ws, row, headers, col_start=1):
    for i, h in enumerate(headers):
        cell = ws.cell(row=row, column=col_start+i, value=h)
        cell.fill      = dark_fill()
        cell.font      = white_bold()
        cell.alignment = center_align()

# ── Section A: Assumptions ───────────────────────────────────
write_section_header(ws4, 1, 1, 4, "\u2699 Assumptions")

# Row 2 empty separator
# Assumption data
assumptions = [
    (3,  "Base EXP to Level",    100,  "EXP",          "SharedExperienceManager.baseExpToLevel"),
    (4,  "EXP Growth Rate",      1.25, "multiplier",    "SharedExperienceManager.expGrowthRate"),
    (5,  "Base Enemy EXP",       10,   "EXP/kill",      "\u0e04\u0e48\u0e32 base \u0e01\u0e48\u0e2d\u0e19 wave scaling"),
    (6,  "Enemy Spawn Rate",     1.5,  "seconds/enemy", "baseSpawnRate \u0e43\u0e19 EnemySpawner"),
    (7,  "Zone Objective Bonus", 1.5,  "multiplier",    "EXP Bonus \u0e40\u0e21\u0e37\u0e48\u0e2d\u0e17\u0e33 zone \u0e2a\u0e33\u0e40\u0e23\u0e47\u0e08"),
    (8,  "Zone Bonus Duration",  30,   "seconds",       "\u0e23\u0e30\u0e22\u0e30\u0e40\u0e27\u0e25\u0e32 EXP bonus"),
    (9,  "Wave EXP Mult/Wave",   0.15, "+% per wave",   "expMultPerWave"),
    (10, "Wave Duration",        60,   "seconds",       "\u0e01\u0e48\u0e2d\u0e19\u0e02\u0e36\u0e49\u0e19 wave \u0e16\u0e31\u0e14\u0e44\u0e1b"),
    (11, "Player Count",         2,    "players",       "\u0e2a\u0e48\u0e07\u0e1c\u0e25\u0e15\u0e48\u0e2d enemy \u0e17\u0e35\u0e48\u0e06\u0e48\u0e32\u0e44\u0e14\u0e49\u0e15\u0e48\u0e2d\u0e19\u0e32\u0e17\u0e35"),
]

for row_num, label, value, unit, notes in assumptions:
    # Col A
    ca = ws4.cell(row=row_num, column=1, value=label)
    ca.font = plain_font(bold=True)
    ca.alignment = left_align()

    # Col B (user input)
    cb = ws4.cell(row=row_num, column=2, value=value)
    cb.font = plain_font(color=BLUE_TEXT)
    cb.fill = PatternFill("solid", fgColor=LIGHT_YELLOW)
    cb.alignment = center_align()

    # Col C
    cc = ws4.cell(row=row_num, column=3, value=unit)
    cc.alignment = left_align()

    # Col D
    cd = ws4.cell(row=row_num, column=4, value=notes)
    cd.alignment = left_align(wrap=True)

# Col widths for section A
ws4.column_dimensions["A"].width = 25
ws4.column_dimensions["B"].width = 12
ws4.column_dimensions["C"].width = 14
ws4.column_dimensions["D"].width = 40

# ── Section B: EXP per Level ─────────────────────────────────
write_section_header(ws4, 13, 1, 4, "\U0001f4ca EXP Required per Level")
write_header_row(ws4, 14, ["Level", "EXP to Next", "Cumulative EXP", "Notes"])

level_notes = {5: "Mini Boss territory", 10: "Mid-game", 15: "Late game"}

for i in range(1, 16):  # Levels 1-15
    row = 14 + i  # rows 15-29
    use_alt = (i % 2 == 0)
    bg_color = ALT_ROW_GRAY if use_alt else "FFFFFF"
    bg_fill  = PatternFill("solid", fgColor=bg_color)

    # Col A: level number
    ca = ws4.cell(row=row, column=1, value=i)
    ca.fill = bg_fill
    ca.alignment = center_align()
    ca.number_format = "#,##0"

    # Col B: EXP to next = INT($B$3 * $B$4^(level-1))
    cb = ws4.cell(row=row, column=2,
                  value=f"=INT($B$3*$B$4^(A{row}-1))")
    cb.fill = bg_fill
    cb.alignment = center_align()
    cb.number_format = "#,##0"

    # Col C: Cumulative
    if i == 1:
        cc_val = 0
        ws4.cell(row=row, column=3, value=0).number_format = "#,##0"
        ws4.cell(row=row, column=3).fill = bg_fill
        ws4.cell(row=row, column=3).alignment = center_align()
    else:
        prev_row = row - 1
        cc = ws4.cell(row=row, column=3,
                      value=f"=C{prev_row}+B{prev_row}")
        cc.fill = bg_fill
        cc.alignment = center_align()
        cc.number_format = "#,##0"

    # Col D: Notes
    cd = ws4.cell(row=row, column=4, value=level_notes.get(i, ""))
    cd.fill = bg_fill
    cd.alignment = left_align()

# Row 30 empty separator (already blank)

# ── Section C: Timeline EXP Model ────────────────────────────
write_section_header(ws4, 31, 1, 8, "\u23f1 Timeline EXP Model (per 30 seconds)")
timeline_headers = [
    "Time", "Minutes", "Wave #", "EXP/min (base)",
    "Zone Bonus Active", "EXP Gained (30s)", "Total EXP", "Est. Level"
]
write_header_row(ws4, 32, timeline_headers)

# Set col widths for section C
col_widths_c = {"A": 8, "B": 10, "C": 10, "D": 16,
                "E": 18, "F": 18, "G": 14, "H": 12}
for col_letter, w in col_widths_c.items():
    ws4.column_dimensions[col_letter].width = w

# Time stamps every 30s from 0:00 to 15:00 = 31 entries
import math

zone_bonus_minutes = {2.0, 2.5, 6.0, 6.5, 8.0, 8.5, 12.0, 12.5}

green_fill = PatternFill("solid", fgColor=GREEN_BG)

for step in range(31):  # 0..30
    minutes = step * 0.5
    row = 33 + step

    total_seconds = int(minutes * 60)
    mins_part = total_seconds // 60
    secs_part = total_seconds % 60
    time_str = f"{mins_part}:{secs_part:02d}"

    is_zone = minutes in zone_bonus_minutes
    row_fill = green_fill if is_zone else PatternFill("solid", fgColor="FFFFFF")

    # Col A: time
    ca = ws4.cell(row=row, column=1, value=time_str)
    ca.fill = row_fill
    ca.alignment = center_align()

    # Col B: minutes
    cb = ws4.cell(row=row, column=2, value=minutes)
    cb.fill = row_fill
    cb.alignment = center_align()
    cb.number_format = "0.0"

    # Col C: Wave # = INT(B{row}/$B$10)+1
    cc = ws4.cell(row=row, column=3,
                  value=f"=INT(B{row}/$B$10)+1")
    cc.fill = row_fill
    cc.alignment = center_align()

    # Col D: EXP/min base = (60/$B$6)*$B$5*$B$11*(1+(C{row}-1)*$B$9)
    cd = ws4.cell(row=row, column=4,
                  value=f"=(60/$B$6)*$B$5*$B$11*(1+(C{row}-1)*$B$9)")
    cd.fill = row_fill
    cd.alignment = center_align()
    cd.number_format = "0.0"

    # Col E: Zone Bonus Active
    ce = ws4.cell(row=row, column=5,
                  value="\u2713" if is_zone else "")
    ce.fill = row_fill
    ce.alignment = center_align()

    # Col F: EXP Gained 30s = IF(E{row}="✓", D{row}/2*$B$7, D{row}/2)
    cf = ws4.cell(row=row, column=6,
                  value=f'=IF(E{row}="\u2713",D{row}/2*$B$7,D{row}/2)')
    cf.fill = row_fill
    cf.alignment = center_align()
    cf.number_format = "0"

    # Col G: Total EXP
    if step == 0:
        cg = ws4.cell(row=row, column=7, value=f"=F{row}")
    else:
        prev = row - 1
        cg = ws4.cell(row=row, column=7, value=f"=G{prev}+F{row}")
    cg.fill = row_fill
    cg.alignment = center_align()
    cg.number_format = "#,##0"

    # Col H: Est. Level = IFERROR(MATCH(G{row},$C$15:$C$29,1),1)
    ch = ws4.cell(row=row, column=8,
                  value=f"=IFERROR(MATCH(G{row},$C$15:$C$29,1),1)")
    ch.fill = row_fill
    ch.alignment = center_align()
    ch.font = plain_font(bold=True, color=GREEN_TEXT)  # default green; below we'll vary

# Apply level-based color to Col H based on estimated level bands
from openpyxl.formatting.rule import CellIsRule

# Green: level 1-5
rule_green = CellIsRule(operator='between', formula=['1', '5'],
                        font=Font(bold=True, color=GREEN_TEXT))
ws4.conditional_formatting.add(f"H33:H{33+30}", rule_green)

# Orange: level 6-10
rule_orange = CellIsRule(operator='between', formula=['6', '10'],
                         font=Font(bold=True, color=ORANGE))
ws4.conditional_formatting.add(f"H33:H{33+30}", rule_orange)

# Red: level 11+
rule_red = CellIsRule(operator='greaterThanOrEqual', formula=['11'],
                      font=Font(bold=True, color=RED))
ws4.conditional_formatting.add(f"H33:H{33+30}", rule_red)

# Note below table row 65
note_cell = ws4.cell(row=65, column=1,
    value="\u26a0 \u0e40\u0e27\u0e25\u0e32\u0e17\u0e35\u0e48\u0e41\u0e2a\u0e14\u0e07\u0e40\u0e1b\u0e47\u0e19\u0e04\u0e48\u0e32\u0e1b\u0e23\u0e30\u0e21\u0e32\u0e13 \u2014 \u0e02\u0e36\u0e49\u0e19\u0e2d\u0e22\u0e39\u0e48\u0e01\u0e31\u0e1a playstyle, \u0e08\u0e33\u0e19\u0e27\u0e19\u0e1c\u0e39\u0e49\u0e40\u0e25\u0e48\u0e19 \u0e41\u0e25\u0e30 upgrade \u0e17\u0e35\u0e48\u0e40\u0e25\u0e37\u0e2d\u0e01")
note_cell.font = plain_font(italic=True, color=GRAY_TEXT)
ws4.merge_cells(start_row=65, start_column=1, end_row=65, end_column=8)
note_cell.alignment = left_align()

# ============================================================
# SHEET 1: Add Column H "Est. Level"
# ============================================================
ws1 = wb["Heartbeat Graph"]

# Header H1
h1 = ws1.cell(row=1, column=8, value="Est. Level")
h1.fill      = dark_fill()
h1.font      = white_bold()
h1.alignment = center_align()

# Col width
ws1.column_dimensions["H"].width = 12

# Data rows 2-32
for row in range(2, 33):
    formula = (
        f"=IFERROR(INDEX('Level Up Estimate'!$H$33:'Level Up Estimate'!$H$63,"
        f"MATCH(B{row},'Level Up Estimate'!$B$33:'Level Up Estimate'!$B$63,0)),\"\")"
    )
    cell = ws1.cell(row=row, column=8, value=formula)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.font = plain_font(bold=True)

# Conditional formatting on H col in sheet 1 (same color bands as sheet 4)
ws1.conditional_formatting.add("H2:H32",
    CellIsRule(operator='between', formula=['1','5'],
               font=Font(bold=True, color=GREEN_TEXT)))
ws1.conditional_formatting.add("H2:H32",
    CellIsRule(operator='between', formula=['6','10'],
               font=Font(bold=True, color=ORANGE)))
ws1.conditional_formatting.add("H2:H32",
    CellIsRule(operator='greaterThanOrEqual', formula=['11'],
               font=Font(bold=True, color=RED)))

# ============================================================
# SAVE
# ============================================================
wb.save(XLSX_PATH)
print("Saved OK:", XLSX_PATH)
print("Sheets:", wb.sheetnames)
